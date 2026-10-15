from __future__ import annotations

import argparse
import csv
import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".jsonl", ".json", ".csv", ".xlsx", ".xls"}


def count_records(path: Path) -> int | None:
    suffix = path.suffix.lower()
    try:
        if suffix == ".jsonl":
            with path.open("r", encoding="utf-8-sig") as stream:
                return sum(1 for line in stream if line.strip())
        if suffix == ".json":
            with path.open("r", encoding="utf-8-sig") as stream:
                data = json.load(stream)
            return len(data) if isinstance(data, list) else 1
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                return max(0, sum(1 for _ in csv.reader(stream)) - 1)
        if suffix == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(path, on_demand=True)
            try:
                return max(0, workbook.sheet_by_index(0).nrows - 1)
            finally:
                workbook.release_resources()
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                worksheet = workbook.active
                return max(0, (worksheet.max_row or 0) - 1)
            finally:
                workbook.close()
    except Exception:
        return None
    return None


def infer_mode(path: Path) -> str:
    name = path.name.lower()
    if "search" in name:
        return "关键词搜索"
    if "detail" in name:
        return "视频详情"
    if "creator" in name:
        return "创作者主页"
    return "采集结果"


def list_files(root: Path) -> dict[str, Any]:
    root = root.expanduser().resolve()
    if not root.exists():
        return {"files": []}
    if not root.is_dir():
        raise ValueError("输出路径不是文件夹")

    files = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        stat = path.stat()
        files.append({
            "id": path.relative_to(root).as_posix(),
            "name": path.name,
            "path": path.relative_to(root).as_posix(),
            "mode": infer_mode(path),
            "type": path.suffix[1:].upper(),
            "records": count_records(path),
            "size": stat.st_size,
            "updated": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M"),
            "preview": [],
        })
    files.sort(key=lambda item: item["id"], reverse=True)
    return {"files": files}


def preview_file(root: Path, relative_path: str, limit: int = 100) -> dict[str, Any]:
    root = root.expanduser().resolve()
    target = (root / relative_path).resolve()
    try:
        target.relative_to(root)
    except ValueError as error:
        raise ValueError("文件路径超出输出目录") from error
    if not target.is_file() or target.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError("文件不存在或格式暂不支持")

    suffix = target.suffix.lower()
    rows: list[dict[str, Any]] = []
    columns: list[str] = []
    total = 0

    if suffix == ".jsonl":
        with target.open("r", encoding="utf-8-sig") as stream:
            for line in stream:
                if not line.strip():
                    continue
                total += 1
                if len(rows) >= limit:
                    continue
                value = json.loads(line)
                rows.append(value if isinstance(value, dict) else {"内容": value})
    elif suffix == ".json":
        with target.open("r", encoding="utf-8-sig") as stream:
            data = json.load(stream)
        values = data if isinstance(data, list) else [data]
        total = len(values)
        rows = [value if isinstance(value, dict) else {"内容": value} for value in values[:limit]]
    elif suffix == ".csv":
        with target.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            columns = list(reader.fieldnames or [])
            for row in reader:
                total += 1
                if len(rows) < limit:
                    rows.append(dict(row))
    elif suffix in {".xlsx", ".xls"}:
        import pandas as pd

        frame = pd.read_excel(target, nrows=limit)
        columns = [str(column) for column in frame.columns]
        rows = frame.astype(object).where(frame.notna(), None).to_dict(orient="records")
        total = count_records(target) or len(rows)

    if not columns and rows:
        columns = list(rows[0].keys())
    return {"data": rows, "total": total, "columns": columns}


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="action", required=True)

    list_parser = subparsers.add_parser("list-files")
    list_parser.add_argument("root")

    preview_parser = subparsers.add_parser("preview")
    preview_parser.add_argument("root")
    preview_parser.add_argument("path")
    preview_parser.add_argument("--limit", type=int, default=100)

    args = parser.parse_args()
    try:
        if args.action == "list-files":
            result = list_files(Path(args.root))
        else:
            result = preview_file(Path(args.root), args.path, max(1, min(args.limit, 500)))
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
